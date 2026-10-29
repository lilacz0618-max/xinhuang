"""把第 8 页的六步流程并入第 7 页右侧空白，删除第 8 页，合并备注，并把页码改为连续的“n / 总页数”。
    python3 merge78.py in.pptx out.pptx  （需能 import 制作包的 ppt/tools/pptkit.py，路径由 PPTKIT 环境变量给出）
"""
import sys, os, re, copy
sys.path.insert(0, os.environ['PPTKIT'])
import pptkit
from pptkit import box, textbox, C, est_height_in, FIT_LOG
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn

src, out = sys.argv[1], sys.argv[2]
prs = Presentation(src)
s7, s8 = prs.slides[6], prs.slides[7]
W = C['white']

# ---- 从第 8 页读出六步内容（汇报人修改后的文字，原样使用） ----
txt = [sh.text_frame.text for sh in s8.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
heads = [t for t in txt if re.match(r'^[①②③④⑤⑥] ', t)]
assert len(heads) == 6, heads
rows = []
for h in heads:
    i = txt.index(h)
    body = txt[i + 1]
    out_t = txt[i + 2] if i + 2 < len(txt) and txt[i + 2].startswith('得到：') else ''
    rows.append((h, body, out_t))
groups = [t.strip() for t in txt if re.match(r'^[一二]、', t)]          # “一、需求是什么”“二、供给怎么配”

# ---- 在第 7 页右侧绘制纵向六步流程 ----
x0, y0, x1, y1 = 6.15, 1.55, 12.83, 6.95
g = 0.1; rh = (y1 - y0 - 5 * g) / 6
gw, lw, ow = 0.42, 1.72, 1.66
lx = x0 + gw + 0.06; bx = lx + lw + 0.06; ox = x1 - ow
cols = [C['grey'], C['teal'], C['teal'], C['orange'], C['orange'], C['navy']]
bgs = [C['light'], C['teal_l'], C['teal_l'], C['orange_l'], C['orange_l'], C['light']]


def fit_box(x, y, w, h, paras, size, minimum, name, **kw):
    tb, sz, over = textbox(s7, x, y, w, h, paras, size=size, minimum=minimum, tag={'slide': 7, 'box': name}, **kw)
    return tb


# 贯穿六步的连线（置于色块之下）
box(s7, lx + lw / 2 - 0.015, y0 + rh / 2, 0.03, 5 * (rh + g), fill=C['navy'])
for i, (h, body, o) in enumerate(rows):
    y = y0 + i * (rh + g)
    box(s7, lx, y, lw, rh, fill=cols[i])
    fit_box(lx + 0.04, y, lw - 0.08, rh, [h], 13, 11, f'step{i + 1}_h', color=W, bold=True, align='c', anchor='m')
    bw = (ox - 0.06 - bx) if o else (x1 - bx)
    box(s7, bx, y, bw, rh, fill=bgs[i])
    fit_box(bx + 0.04, y, bw - 0.08, rh, [body], 12, 10.5, f'step{i + 1}', anchor='m')
    if o:
        box(s7, ox, y, ow, rh, fill=W, line=cols[i])
        fit_box(ox + 0.03, y, ow - 0.06, rh, [o], 11.5, 10, f'step{i + 1}_out', color=cols[i], bold=True,
                align='c', anchor='m')
# “需求是什么”“供给怎么配”竖排分组条
for k, (r0, col) in enumerate(((1, C['teal']), (3, C['orange']))):
    y = y0 + r0 * (rh + g); h = 2 * rh + g
    lab = re.sub(r'^[一二]、', '', groups[k]) if k < len(groups) else ['需求是什么', '供给怎么配'][k]
    box(s7, x0, y, gw, h, fill=col)
    fit_box(x0, y, gw, h, list(lab), 12, 11, f'group{k + 1}', color=W, bold=True, align='c', anchor='m',
            space_after=0, line=1.0)

# ---- 合并备注 ----
n7 = s7.notes_slide.notes_text_frame.text
n8 = s8.notes_slide.notes_text_frame.text
talk = ('这是技术路线。左边是完整的技术路线图，右边把它简化成六步：准备数据、找规律、做预测、排固定班、派机动车、统一评价。'
        '上一步的结果就是下一步的输入：还原出的真实需求既用来算权重，也用来做预测；预测又直接决定机动车派到哪里。'
        '为了避免看到结果再改方法，计算前做了预注册；实际计算中有四处改动，每一处都标明“偏离预注册”。')
rest = n8.split('\n\n', 1)[1] if '\n\n' in n8 else n8
s7.notes_slide.notes_text_frame.text = talk + '\n\n' + rest

# ---- 删除第 8 页 ----
lst = prs.slides._sldIdLst
sid = list(lst)[7]
prs.part.drop_rel(sid.get(qn('r:id')))
lst.remove(sid)

# ---- 页码改为连续编号 ----
total = len(prs.slides)
for i, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if sh.has_text_frame and re.fullmatch(r'\s*\d+\s*/\s*\d+\s*', sh.text_frame.text):
            runs = sh.text_frame.paragraphs[0].runs
            runs[0].text = f'{i} / {total}'
            for r in runs[1:]:
                r.text = ''
prs.save(out)
over = [d for d in FIT_LOG if d['overflow']]
print(out, total, '页；第 7 页新增文本框', len(FIT_LOG), '个，溢出', len(over), over)
