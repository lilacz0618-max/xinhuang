"""第四轮构建脚本：从 ppt/ppt_spec.json（由 ppt/make_spec_r4.py 生成）生成《新晃城乡公交研究进展汇报》20 页中文 PPT。
第四轮：以第一轮版式为底稿，改写第 3、6、7、9、10、12、14、15、17、18 页（通俗表述、模拟方法与结果摘要、“十五五”规划对接），
新增文献“十五五”规划并重编号。页面文字按第一轮编号书写，写入幻灯片时由 cite_remap 统一换成新编号。


    cd v8 && python build_ppt.py [输出文件名]

在 ppt/tools/build_ppt_starter.py 的基础上改写：
  1) 第 3、4、5、6、10、12、14、15、17、18 页的 diagram 区域全部改为原生形状图示（色块、箭头、卡片、格子、时间轴）；
  2) 上屏要点按版面压缩（不改事实、数字、文献编号与术语），完整原文由 set_notes(..., bullets=...) 写入备注末尾“完整要点：”；
  3) 结果页插图左对齐，右侧文字区随实际图宽加宽；
  4) 所有文字框都经 fit_size 拟合并写入 ppt/fit_log.json，供 self_check.py 检查溢出。
"""
import sys, os, re, json
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ppt', 'tools'))
from pptkit import *
import pptkit

os.chdir(V8)
spec = json.load(open('ppt/ppt_spec.json', encoding='utf-8'))                  # 新编号：标题、所属部分、备注
spec_old = json.load(open('ppt/ppt_spec_r4_oldnum.json', encoding='utf-8'))    # 第一轮编号：页面文字
OUT = sys.argv[1] if len(sys.argv) > 1 else '新晃城乡公交研究进展汇报.pptx'
TOTAL, SYN = len(spec['slides']), spec['synthetic_footer']
SP = {sp['no']: sp for sp in spec_old['slides']}
W = C['white']
NO = [0]                                   # 当前页号（供拟合日志使用）
REMAP = [True]                             # 参考文献页关闭编号映射


def cite_remap(t):
    """第一轮编号 → 第四轮编号：“十五五”规划（记为 [50]）按首次出现排为 [3]，原 [3]—[49] 顺延一位。"""
    f = lambda n: 3 if n == 50 else n if n <= 2 else n + 1
    return re.sub(r'\[(\d+(?:[,，\-–—]\d+)*)\]',
                  lambda m: '[' + re.sub(r'\d+', lambda k: str(f(int(k.group()))), m.group(1)) + ']', t)


_add_rich = pptkit.add_rich
pptkit.add_rich = lambda para, text, *a, **k: _add_rich(para, cite_remap(text) if REMAP[0] else text, *a, **k)


# ---------- 本脚本的小工具 ----------
def T(s, x, y, w, h, paras, size=16, minimum=None, name='text', **kw):
    """带拟合日志的文本框。"""
    return textbox(s, x, y, w, h, paras, size=size, minimum=minimum or size,
                   tag={'slide': NO[0], 'box': name}, **kw)


def lbox(s, x, y, w, h, paras, fill=C['light'], line=None, color=C['text'], size=14, minimum=None,
         bold=False, align='c', anchor='m', shape='rect', name='box', dash=False, bullet=False, inset=0.0):
    """色块 + 拟合文字（文字单独成框，便于检查溢出，也便于汇报人修改）。"""
    box(s, x, y, w, h, fill=fill, line=line, shape=shape, dash=dash)
    return T(s, x + inset, y, w - 2 * inset, h, paras, size, minimum, name, color=color, bold=bold,
             align=align, anchor=anchor, bullet=bullet)


def head_card(s, x, y, w, h, head, paras, col, bg, size=16, minimum=None, head_h=0.45, head_size=16,
              name='card', bullet=True, head_color=W, anchor='t'):
    """标题条 + 正文卡片。"""
    box(s, x, y, w, h, fill=bg)
    box(s, x, y, w, head_h, fill=col, text=head, size=head_size, color=head_color, bold=True, align='l')
    return T(s, x + 0.06, y + head_h + 0.06, w - 0.12, h - head_h - 0.1, paras, size, minimum, name,
             bullet=bullet, anchor=anchor)


def fig(s, path, x, y, w, h, caption, halign='l'):
    """等比插图（默认左对齐）+ 图注，返回图片右边界。"""
    pic, (px, py, pw, ph) = picture_fit(s, path, x, y, w, h - 0.3, valign='t', halign=halign)
    textbox(s, px, py + ph + 0.02, max(pw, 3.0), 0.28, [caption], size=SIZE['caption'], color=C['grey'], fit=False)
    return px + pw, py + ph + 0.3


def caption_of(sp, fo):
    base = os.path.basename(fo['full'])[2:-4]
    return '图：' + base + ('（真实 GIS 底图与现行线路）' if fo['real_basemap'] else
                         '（模拟数据，SYNTHETIC）' if sp['synthetic'] else '')


def hyp_box(s, sp, x, y, w, h):
    lbox(s, x, y, w, h, [sp['hyp_box']], fill=C['teal_l'], line=C['teal'], color=C['teal'], size=16,
         bold=True, name='hyp')


def section_pages(item):
    p = [sl['no'] for sl in spec['slides'] if sl['section'] == item]
    return '' if not p else f'第 {min(p)} 页' if min(p) == max(p) else f'第 {min(p)}—{max(p)} 页'


# ---------- 逐页 ----------
def p1(s, sp):
    t1, _, t2 = sp['title'].partition('——')
    box(s, 0, 0, SW, 0.22, fill=C['navy'])
    k = t1.index('供需')                       # 按语义断行：研究对象 / 研究内容 / 案例地
    T(s, 0.9, 1.1, 11.5, 2.6, [t1[:k], t1[k:], '——' + t2], 34, 28, 'title', color=C['navy'], bold=True, anchor='b',
      space_after=0.1)
    box(s, 0.9, 3.82, 1.2, 0.05, fill=C['teal'])
    T(s, 0.9, 3.95, 11.5, 0.65, [sp['bullets'][0]], 22, 22, 'subtitle', color=C['teal'], bold=True)
    T(s, 0.9, 4.65, 11.5, 1.2, sp['bullets'][1:3], 16, 16, 'info', color=C['grey'])
    lbox(s, 0.9, 6.45, 11.5, 0.5, [sp['bullets'][3]], fill=C['orange_l'], color=C['orange'], size=11,
         align='l', name='synthetic')


def p2(s, sp):
    x, y0, w, h = 1.6, 1.7, 10.1, 5.1
    n = len(sp['bullets']); hh = (h - 0.15 * (n - 1)) / n
    for i, item in enumerate(sp['bullets']):
        y = y0 + i * (hh + 0.15)
        col = C['teal'] if i < 3 else C['orange'] if i == 4 else C['navy']
        box(s, x, y, 0.75, hh, fill=col, text=str(i + 1), size=20, color=W, bold=True)
        lbox(s, x + 0.85, y, w - 0.85, hh, [item], color=C['navy'], size=20, bold=True, align='l', inset=0.1,
             name=f'toc{i + 1}')
        textbox(s, x + w - 2.2, y, 2.1, hh, [section_pages(item)], size=12, color=C['grey'], align='r',
                anchor='m', fit=False)


def p3(s, sp):
    # 左栏：三个纵向色块
    items = [('政策要求', C['navy'],
              '2011 年起推进城乡道路客运一体化，目标从“有车坐”走向“坐得上、坐得好”[1]；2024 年服务指南要求按客流规律灵活安排班次，鼓励预约响应[2]'),
             ('山区现实', C['teal'],
              '交通劣势与相对贫困存在可测的关联[3]；医疗等公共服务的可达性明显受交通条件限制[4]'),
             ('理论视角', C['grey'],
              '缺少公交使农村居民难以到达就业、医疗和教育机会，属于社会排斥问题[5,6]；城乡交通融合已成为地理学研究议题[7]')]
    bh = (5.3 - 0.24) / 3
    for i, (lab, col, txt) in enumerate(items):
        y = 1.6 + i * (bh + 0.12)
        box(s, 0.5, y, 6.05, bh, fill=C['light'])
        box(s, 0.5, y, 1.3, bh, fill=col, text=lab, size=17, color=W, bold=True)
        T(s, 1.9, y + 0.04, 4.6, bh - 0.08, [txt], 16, 15, f'left{i + 1}', anchor='m')
    # 右栏：研究对象卡片（含真实线网底图）
    x, y, w, h = 6.78, 1.6, 6.05, 5.3
    box(s, x, y, w, h, fill=C['light'])
    box(s, x, y, w, 0.45, fill=C['navy'], text='研究对象', size=16, color=W, bold=True, align='l')
    T(s, x + 0.15, y + 0.52, w - 0.3, 0.56, ['湖南省新晃侗族自治县（武陵山区）'], 20, 18, 'obj1', color=C['navy'], bold=True)
    T(s, x + 0.15, y + 1.08, w - 0.3, 0.42, ['县汽车站发出的 7 条城乡公交线路'], 16, 16, 'obj2', color=C['teal'], bold=True)
    fo = SP[8]['figure_options'][0]
    fig(s, fo['full'], x + 0.15, y + 1.55, w - 0.3, h - 1.6, '图：研究区与城乡公交线网（真实 GIS 底图与现行线路，详见第 8 页）',
        halign='c')


def p4(s, sp):
    box(s, 0.5, 1.6, 12.33, 1.25, fill=C['light'])
    box(s, 0.5, 1.6, 2.0, 1.25, fill=C['navy'], text=['理论依据', '时间维度的可达性'], size=15, color=W, bold=True)
    T(s, 2.6, 1.6, 10.15, 1.25, ['人的活动受能力、组合和权威三类约束，必须在特定时间到达特定地点[8]；时空可达性是交通地理的前沿议题[9,10]；'
                                 '公交可达性在一天之内随时刻表剧烈波动[11]。农村乘客以老人、学生和无车者为主[12]，班次时刻直接决定其能否出行。'],
      16, 15, 'banner', anchor='m')
    cards = [('难题一', '权重', '班次依据什么？', '影响因子的权重多用熵权法、灰色关联给出，难以回答“客流变化由什么造成”', '对应 RQ1', C['teal'], C['teal_l']),
             ('难题二', '删失', '台账能否代表需求？', '满载班次的台账只记录“坐上车的人”，真实需求被座位数截住，形成删失', '对应 RQ1、RQ2', C['teal'], C['teal_l']),
             ('难题三', '分工', '固定班与机动车如何分工？', '固定时刻表优化与机动加班分开讨论，缺少以预测为纽带的分工', '对应 RQ3、RQ4', C['orange'], C['orange_l'])]
    cw, gap = 3.98, (12.33 - 3 * 3.98) / 2
    for i, (no, word, q, body, rq, col, bg) in enumerate(cards):
        x, y, h = 0.5 + i * (cw + gap), 3.05, 3.85
        box(s, x, y, cw, h, fill=C['white'], line=C['silver'])
        box(s, x, y, cw, 0.08, fill=col)
        textbox(s, x + 0.2, y + 0.22, 1.2, 0.3, [no], size=12, color=C['grey'], fit=False)
        T(s, x + 0.2, y + 0.42, cw - 0.4, 0.8, [word], 30, 30, f'card{i + 1}_word', color=col, bold=True)
        T(s, x + 0.2, y + 1.22, cw - 0.4, 0.45, [q], 17, 16, f'card{i + 1}_q', color=C['navy'], bold=True)
        box(s, x + 0.2, y + 1.72, cw - 0.4, 0.02, fill=C['silver'])
        T(s, x + 0.2, y + 1.82, cw - 0.4, 1.3, [body], 16, 16, f'card{i + 1}_body')
        lbox(s, x + 0.2, y + h - 0.6, 2.4, 0.4, ['→ ' + rq], fill=bg, color=col, size=13, bold=True,
             name=f'card{i + 1}_rq')


def p5(s, sp):
    labels = ['需求的时间结构', '权重识别', '删失还原', '概率预测', '两层供给与预测衔接']
    cols = [C['teal']] * 4 + [C['orange']]
    segs = chain(s, labels, 0.5, 1.6, 12.33, 0.75, colors=cols, size=15)
    txt = ['个人出行在日与日之间既有习惯也有变异[13]；定期集市是农村社会经济结构的基本单元[14]，其对公交客流的影响需数据检验',
           '常用熵权法、灰色关联等客观赋权[15]，按离散程度或序列相似程度赋权，不衡量对客流变化的解释力',
           '收益管理早已指出：售罄后的销售量会低估真实需求，需要“需求还原”[16]',
           '国内研究多针对城市地铁和公交的短时客流[17,18]，县域班次级的小计数预测较少',
           '半柔性公交：固定骨架 + 有限柔性[19]；农村需求响应调度刚受关注[20]；两步法忽略预测误差[21]']
    for i, ((x, w), t) in enumerate(zip(segs, txt)):
        T(s, x + 0.02, 2.45, w - 0.04, 2.0, [t], 14, 13, f'chain{i + 1}')
    T(s, 0.5, 4.55, 6.0, 0.4, ['四个相互关联的缺口'], 16, 16, 'gaps_head', color=C['navy'], bold=True)
    gaps = ['① 多重节律的相对权重缺少同时估计', '② 删失需求在城乡公交中没有被处理',
            '③ 县域班次级概率预测缺位', '④ 固定层与柔性层缺少以预测为纽带的协同设计']
    rqs = [('RQ1', C['teal']), ('RQ1、RQ2', C['teal']), ('RQ2', C['teal']), ('RQ3、RQ4', C['orange'])]
    gw, gg = (12.33 - 3 * 0.35) / 4, 0.35
    for i, g in enumerate(gaps):
        x = 0.5 + i * (gw + gg)
        lbox(s, x, 4.98, gw, 1.05, [g], fill=C['white'], line=C['navy'], color=C['navy'], size=15, bold=True,
             name=f'gap{i + 1}', inset=0.05)
        if i < 3:
            arrow(s, x + gw + 0.04, 5.5, x + gw + gg - 0.04, 5.5, color=C['navy'])
        arrow(s, x + gw / 2, 6.05, x + gw / 2, 6.38, color=rqs[i][1])
        lbox(s, x + gw / 2 - 0.8, 6.42, 1.6, 0.45, [rqs[i][0]], fill=rqs[i][1], color=W, size=14, bold=True,
             shape='round', name=f'gap{i + 1}_rq')


def p6(s, sp):
    mx, my, mw, mh = 0.5, 1.6, 7.7, 4.05
    rq = [('RQ1　节律权重', '时段、周节律、校历、节假日、春运、集日、天气、季节八个因子族对班次客流的相对影响各有多大？集日是否是主要因素？'),
          ('RQ2　删失与预测', '存在满载删失时，什么模型能给出班次级次日预测，并在春运、节假日和天气事件日保持可靠的预测区间？'),
          ('RQ3　固定层', '只承担稳定节律的固定时刻表应如何编制？各线路的固定发车数量应为多少？'),
          ('RQ4　柔性层', '少量机动车按次日预测指派，能在多大程度上降低社会成本？收益集中在哪些日期？预测改进还有多大价值？')]
    lw, gap = 0.55, 0.1
    cw, ch = (mw - lw - 2 * gap) / 2, (mh - gap) / 2
    for r, (lab, col, bg) in enumerate([('需求', C['teal'], C['teal_l']), ('供给', C['orange'], C['orange_l'])]):
        y = my + r * (ch + gap)
        box(s, mx, y, lw, ch, fill=col, text=list(lab), size=18, color=W, bold=True)
        for c in range(2):
            k = r * 2 + c; x = mx + lw + gap + c * (cw + gap)
            head_card(s, x, y, cw, ch, rq[k][0], [rq[k][1]], col, bg, size=16, minimum=15, bullet=False,
                      name=f'rq{k + 1}')
    # 假设栏
    hyp = [('H1', 'RQ1', '日内时段是班次层最大的因子族，春运和节假日是日层最大的因子族，集日和天气权重较小'),
           ('H2', 'RQ1', '熵权法、灰色关联的权重排序与真实排序明显不一致'),
           ('H3', 'RQ2', '考虑删失的模型在满载班次上的偏差更小'),
           ('H4', 'RQ2', '全局共形区间在高需求情景覆盖不足，Mondrian 区间接近名义覆盖'),
           ('H5', 'RQ3', '只按稳定节律编制的固定时刻表能减少总班次并降低社会成本'),
           ('H6', 'RQ4', '柔性层收益集中在春运和节假日边沿，普通日接近零')]
    hx, hw, gy = 8.45, 4.38, 0.07
    hh = (5.3 - 5 * gy) / 6
    for i, (h, r, t) in enumerate(hyp):
        y = 1.6 + i * (hh + gy); col = C['teal'] if i < 4 else C['orange']
        box(s, hx, y, hw, hh, fill=C['teal_l'] if i < 4 else C['orange_l'])
        box(s, hx, y, 0.62, hh, fill=col, text=[h, r], size=12, color=W, bold=True)
        T(s, hx + 0.66, y, hw - 0.7, hh, [t], 14, 12.5, f'hyp{i + 1}', anchor='m')
    # 论文横条
    T(s, mx, 5.7, mw, 0.3, ['两篇论文合起来构成学位论文'], 12, 12, 'papers_head', color=C['navy'], bold=True, align='c')
    box(s, mx, 6.0, mw, 0.03, fill=C['navy'])
    pw = (mw - 0.1) / 2
    lbox(s, mx, 6.08, pw, 0.82, ['论文一（RQ1—RQ2）', '删失感知的多重节律识别与概率预测'], fill=C['teal'], color=W,
         size=13, bold=True, name='paper1')
    lbox(s, mx + pw + 0.1, 6.08, pw, 0.82, ['论文二（RQ3—RQ4）', '预测驱动的县域城乡公交半柔性配置'], fill=C['orange'],
         color=W, size=13, bold=True, name='paper2')


def p7(s, sp):
    fo = sp['figure_options'][0]
    right, _ = fig(s, fo['full'], 0.5, 1.5, 7.45, 5.45, caption_of(sp, fo))
    x = right + 0.35; w = 12.83 - x
    n = len(sp['tags']); tw = (w - 0.12 * (n - 1)) / n
    for i, t in enumerate(sp['tags']):
        lbox(s, x + i * (tw + 0.12), 1.5, tw, 0.6, [t], fill=C['navy'], color=W, size=14, bold=True, name=f'tag{i + 1}')
    T(s, x, 2.3, w, 4.6,
      ['四个模块首尾相接：删失负二项同时服务 RQ1 的权重识别和 RQ2 的预测结构项；RQ2 的预测分布驱动 RQ4 的柔性层；RQ3 只承担稳定节律',
       '预注册：计算前写下主方法、对照方法、检验期和评价指标[22]；检验期 2026-01-01 至 2026-09-30，2026-10-01 以后的台账封存，本轮不读取',
       '训练与检验按时间顺序切分，避免把未来信息泄露给模型[23,24]',
       '四处偏离预注册（改进预测方法、最少服务规则、柔性层随机规划求解方式、增加反应式对照），均在正文标注“偏离预注册”并说明理由'],
      17, 16, 'right', bullet=True)


def p8(s, sp):
    fo = sp['figure_options'][0]
    right, bottom = fig(s, fo['full'], 0.5, 1.5, 6.3, 3.9, caption_of(sp, fo))
    lbox(s, 0.5, bottom + 0.1, right - 0.5, 6.95 - bottom - 0.1,
         ['集日与天气只作对照因子：县城集期“逢农历四、九”属待确认信息，需逐乡核实[14]；天气对公交客流的影响因地区和人群而异[25,26]',
          '不纳入模型：工业园区倒班、医院门诊排班、务工人员个人返乡日期（无法可靠获取，影响经节假日、春运和周节律间接体现）'],
         size=14, minimum=13, align='l', anchor='m', bullet=True, inset=0.06, name='left_note')
    x = right + 0.25; w = 12.83 - x
    head_card(s, x, 1.5, w, 2.5, '真实可得的数据',
              ['县汽车站逐日台账（2024-06-01 至 2026-07-27，全县客流总量锚点）', '7 条线路双向现行时刻表',
               '法定节假日与调休、农历、春运、校历与住校生离返校日', '沿线逐日天气实况与预报存档（检验期用提前 1 天的预报值）',
               '行政村点位与人口栅格'], C['teal'], C['teal_l'], size=14, minimum=13, head_h=0.4, name='real')
    head_card(s, x, 4.1, w, 2.2, '需要调查的数据',
              ['暑期跟车调查（已完成，原始问卷待录入）', '寒假与春运跟车、行政村入户、线上问卷（计划），未成年人由家长回答',
               '问卷补充乘客群体构成与期望出发时刻；农村出行受收入、距离和服务水平共同影响[27]'],
              C['orange'], C['orange_l'], size=14, minimum=13, head_h=0.4, name='survey')
    lbox(s, x, 6.4, w, 0.55, ['观测口径：每趟班次只有全程总登乘人数，无站点上下客和 OD'], fill='E9ECEF',
         color=C['navy'], size=14, minimum=13, bold=True, align='l', inset=0.08, name='scope')


def p9(s, sp):
    fo = sp['figure_options'][0]
    right, _ = fig(s, fo['full'], 0.5, 1.5, 7.45, 5.45, caption_of(sp, fo))
    x = right + 0.35; w = 12.83 - x
    T(s, x, 1.5, w, 3.85,
      ['为什么要模拟：真实数据中看不到潜在需求，无法判断权重是否识别正确、删失是否还原到位；模拟系统为每个班次保留“真值”（潜在需求、群体构成、未满足人数）',
       '规模：852 天，116127 个班次观测；满载删失班次占 3.0%；调度员临时加班 255 次；四套模拟问卷共 808 份',
       '用真实的时刻表、日历校历与天气驱动，并与真实县站台账核对总量',
       '真实数据到位后：跳过模拟的两个步骤，按同一流程重算（附录 A 操作指南）'],
      16, 16, 'right', bullet=True, emph=['852', '116127', '3.0%', '255', '808'])
    n = len(sp['number_bar']); cw = (w - 0.1 * (n - 1)) / n
    for i, d in enumerate(sp['number_bar']):
        bx = x + i * (cw + 0.1)
        box(s, bx, 5.5, cw, 1.4, fill=C['light'])
        T(s, bx, 5.55, cw, 0.75, [d['value'] + d['unit']], 22, 16, 'number_' + d['key'], color=C['orange'],
          bold=True, align='c', anchor='m')
        T(s, bx, 6.3, cw, 0.55, [d['label']], 11, 10, 'label_' + d['key'], color=C['grey'], align='c', anchor='t')


def p10(s, sp):
    flow = [(['逐班次台账', '每趟班次的全程总登乘人数'], C['light'], C['navy']),
            (['删失负二项', '满载班次只知 Y* ≥ 座位数 c'], C['teal'], W),
            (['期望需求 log μ̂', '线路×方向基线 + 八个因子族'], C['teal_l'], C['navy']),
            (['因子族 Shapley 分解', 'R² 在 8 个因子族的全部排列上平均分配'], C['teal'], W),
            (['权重与 95% 区间', '以连续 7 天为块的自助法'], C['navy'], W)]
    side = {1: '负二项描述过离散的计数[28]；删失借用收益管理的需求还原思路[29]',
            3: '结果可加，并能处理因子之间的相关[30,31]',
            4: '两个层面：班次层（同一天内班次之间）与日层（日期之间）'}
    bw, bh, g = 3.75, 0.78, 0.35
    for i, (t, fill, col) in enumerate(flow):
        y = 1.6 + i * (bh + g)
        lbox(s, 0.5, y, bw, bh, t, fill=fill, color=col, size=15, minimum=13, bold=True, name=f'flow{i + 1}',
             line=C['teal'] if fill == C['teal_l'] else None)
        if i < 4:
            arrow(s, 0.5 + bw / 2, y + bh + 0.02, 0.5 + bw / 2, y + bh + g - 0.02, color=C['navy'], width=2)
        if i in side:
            box(s, 0.5 + bw + 0.12, y + 0.1, 0.03, bh - 0.2, fill=C['grey'])
            T(s, 0.5 + bw + 0.2, y - 0.12, 6.5 - bw - 0.2 - 0.5, bh + 0.24, [side[i]], 13, 12, f'side{i + 1}',
              color=C['grey'], anchor='m')
    # 公式框
    x, w = 6.75, 6.08
    head_card(s, x, 1.6, w, 1.65, '① 计数模型与删失似然',
              ['log μ = 线路×方向基线 + 时段 + 周节律 + 校历 + 节假日 + 春运 + 集日 + 天气 + 季节',
               '满载班次：Y* ≥ c，似然贡献改写为 P(Y* ≥ c)'], C['navy'], C['light'], size=15, minimum=14,
              head_h=0.4, head_size=15, bullet=False, name='formula1')
    head_card(s, x, 3.37, w, 1.2, '② 因子族 Shapley 权重（K = 8）',
              ['φk = Σ［|S|!(K−|S|−1)!/K!］·［R²(S∪{k}) − R²(S)］'], C['navy'], C['light'], size=16, minimum=14,
              head_h=0.4, head_size=15, bullet=False, name='formula2', anchor='m')
    head_card(s, x, 4.7, w, 2.2, '对照方法',
              ['熵权法、灰色关联（客观赋权）', 'LightGBM-SHAP：特征相关时，SHAP 的分配受依赖结构影响[32]',
               '忽略删失的朴素负二项'], C['grey'], C['light'], size=16, minimum=15, head_h=0.4, head_size=15,
              name='right')


def p11(s, sp):
    fo = sp['figure_options'][0]
    right, _ = fig(s, fo['full'], 0.5, 1.5, 7.45, 5.45, caption_of(sp, fo))
    x = right + 0.35; w = 12.83 - x
    T(s, x, 1.5, w, 4.2,
      ['班次层：时段族权重 83.5%（95% 区间 79.1%—85.9%），模拟真值 85.8%',
       '日层：春运 36.7%、节假日 29.1%；集日 4.5%、天气 3.8%',
       '方法比较（与真值权重的平均绝对误差，百分点）：删失负二项 0.6，灰色关联 9.0，SHAP 10.2，熵权法 10.7；与真值的秩相关：熵权法仅 0.36，灰色关联 0.29',
       '含义：固定时刻表首先要排对“一天之内的班次分布”；集日不宜作为排班的核心依据，这与以集市为中心的农村研究传统[14]形成对话，需用真实数据复核'],
      16, 16, 'right', bullet=True, emph=['83.5%', '36.7%', '29.1%', '0.6'])
    hyp_box(s, sp, x, 5.85, w, 1.05)


def p12(s, sp):
    steps = [('① 结构项', C['teal'], ['删失负二项给出每个班次的期望需求 μ̂']),
             ('② 残差修正', C['teal'], ['LightGBM 学习结构项未捕捉的非线性部分[33]；删失班次的学习目标用 E[Y* | Y* ≥ c] 插补']),
             ('③ 预测区间', C['navy'], ['训练期按月 5 折交叉拟合得样本外残差；共形预测在很弱的假设下给出有覆盖率保证的区间[34,35]'])]
    bw = (12.33 - 2 * 0.45) / 3
    for i, (h, col, t) in enumerate(steps):
        x = 0.5 + i * (bw + 0.45)
        head_card(s, x, 1.6, bw, 1.8, h, t, col, C['light'], size=15, minimum=14, head_h=0.42, bullet=False,
                  name=f'flow{i + 1}')
        if i < 2:
            arrow(s, x + bw + 0.05, 2.5, x + bw + 0.4, 2.5, color=C['navy'], width=2.25)
    # 左：情景划分
    T(s, 0.5, 3.6, 6.05, 0.4, ['情景划分：按 5 个情景分别校准区间（Mondrian）'], 16, 15, 'left_head', color=C['navy'], bold=True)
    sc = [(['春运', '核心'], C['orange']), (['节假日', '边沿'], C['orange']), (['学生', '离返校日'], C['teal']),
          (['天气', '事件日'], C['grey']), (['普通日'], C['navy'])]
    sw = (6.05 - 4 * 0.08) / 5
    for i, (t, col) in enumerate(sc):
        lbox(s, 0.5 + i * (sw + 0.08), 4.08, sw, 0.9, t, fill=col, color=W, size=15, minimum=13, bold=True,
             name=f'scen{i + 1}', inset=0.03)
    T(s, 0.5, 5.05, 6.05, 0.75, ['春运和节假日的波动远大于普通日，统一的区间宽度会使高需求日覆盖不足，因此分情景校准'], 14, 13,
      'scen_note', color=C['grey'])
    lbox(s, 0.5, 5.85, 6.05, 1.05,
         ['改进方法（偏离预注册）：删失感知 LightGBM，删失班次的目标值用 EM 算法迭代插补，单独报告'],
         fill=C['white'], line=C['orange'], dash=True, color=C['text'], size=16, align='l', inset=0.12, name='improved')
    # 右：评价与对照
    head_card(s, 6.78, 3.6, 6.05, 1.9, '评价指标',
              ['MAE、RMSE、WAPE 与平均偏差[36]；区间看覆盖率与平均宽度', 'Diebold–Mariano 检验判断精度差异是否显著[37]',
               '计数预测的专门评价方法[38]'], C['navy'], C['light'], size=16, minimum=14, head_h=0.4, head_size=15,
              name='right_eval')
    head_card(s, 6.78, 5.58, 6.05, 1.32, '对照方法',
              ['季节朴素、历史均值、朴素负二项、删失负二项、LightGBM（泊松）'], C['grey'], C['light'], size=16,
              minimum=15, head_h=0.4, head_size=15, bullet=False, name='right_base')


def p13(s, sp):
    fo = sp['figure_options'][0]
    right, _ = fig(s, fo['full'], 0.5, 1.5, 7.45, 5.45, caption_of(sp, fo))
    x = right + 0.3; w = 12.83 - x
    T(s, x, 1.5, w, 4.3,
      ['总体 WAPE：主方法 42.1%，季节朴素 57.9%，删失负二项 51.1%，LightGBM（泊松）40.0%，改进方法 41.1%',
       '显著性：主方法显著优于季节朴素、历史均值和两种负二项；与两种 LightGBM 差异在 5% 水平上不显著（DM 检验 p = 0.067）',
       '满载删失班次：忽略删失的 LightGBM 平均低估 18.6 人，主方法 9.3 人',
       '90% 区间：春运核心期 Mondrian 覆盖率 92.1%，全局区间 87.6%；Mondrian 各情景覆盖率 90.0%—92.2%',
       '使用建议：两种方法并行，点预测用改进方法，区间用主方法'],
      16, 16, 'right', bullet=True, emph=['42.1%', '18.6', '9.3', '92.1%', '87.6%'])
    hyp_box(s, sp, x, 5.95, w, 0.95)


def p14(s, sp):
    # 左栏：分工 → 需求分层 → 时段反演 → 曲线示意
    lw = 6.05
    lbox(s, 0.5, 1.6, (lw - 0.1) / 2, 0.95, ['固定时刻表：日内时段、周节律、学期/假期交替'], fill=C['navy'], color=W,
         size=14, bold=True, name='split_fixed', inset=0.06)
    lbox(s, 0.5 + (lw + 0.1) / 2, 1.6, (lw - 0.1) / 2, 0.95, ['次日柔性层：春运、节假日、天气等不规则节律'],
         fill=C['orange_l'], line=C['orange'], color=C['orange'], size=14, bold=True, name='split_flex', inset=0.06)
    T(s, 0.5, 2.65, lw, 0.38, ['需求分层：6 层场景'], 15, 15, 'layers_head', color=C['navy'], bold=True)
    layers = ['学期周一至周四', '学期周五', '学期周六', '学期周日', '假期工作日', '假期周末']
    cw = (lw - 2 * 0.08) / 3
    for i, t in enumerate(layers):
        x = 0.5 + (i % 3) * (cw + 0.08); y = 3.05 + (i // 3) * 0.52
        lbox(s, x, y, cw, 0.45, [t], fill=C['teal_l'] if i < 4 else C['light'], line=C['teal'] if i < 4 else C['silver'],
             color=C['navy'], size=14, name=f'layer{i + 1}')
    sw = (lw - 0.4) / 2
    lbox(s, 0.5, 4.15, sw, 0.8, ['逐班次拟合删失负二项', '→ 潜在需求'], fill=C['teal'], color=W, size=13, bold=True, name='inv1')
    arrow(s, 0.5 + sw + 0.05, 4.56, 0.5 + sw + 0.35, 4.56, color=C['navy'], width=2)
    lbox(s, 0.5 + sw + 0.4, 4.15, sw, 0.8, ['问卷期望出发时刻', '→ 分摊到 15 分钟时段'], fill=C['teal'], color=W, size=13,
         bold=True, name='inv2')
    # 15 分钟时段需求曲线（示意形状，非数据）
    import math
    bx, by, bw, bh = 0.7, 5.15, lw - 0.4, 1.35
    box(s, 0.5, 5.05, lw, 1.85, fill=C['light'])
    n = 48
    for k in range(n):
        t = k / (n - 1)
        v = 0.15 + 0.85 * math.exp(-((t - 0.12) / 0.07) ** 2) + 0.55 * math.exp(-((t - 0.62) / 0.12) ** 2)
        hh = bh * min(v, 1.0)
        box(s, bx + k * bw / n, by + bh - hh, bw / n * 0.8, hh, fill=C['teal'] if v > 0.6 else C['silver'])
    box(s, bx, by + bh, bw, 0.015, fill=C['grey'])
    textbox(s, 0.6, 6.55, lw - 0.2, 0.3, ['15 分钟时段需求曲线（示意）：固定发车按此曲线在网格上排布'], size=11,
            color=C['grey'], fit=False)
    # 右栏：公式 + 约束 + 方法来源
    x, w = 6.78, 6.05
    head_card(s, x, 1.6, w, 1.75, '目标：社会成本最小[39,40]',
              ['社会成本 = 运营成本 + 时间价值 × 等待人时 + 未满足惩罚 ×（时刻表抑制损失 + 满载未满足）'],
              C['orange'], C['orange_l'], size=16, minimum=15, head_h=0.42, bullet=False, name='formula', anchor='m')
    head_card(s, x, 3.5, w, 2.6, '决策与约束',
              ['决策：每条线路每个方向在 10 分钟网格上是否发车',
               '约束：乘客分配与容量、车辆平衡、驾驶时长、最少服务（偏离预注册：以现行首末班为界）、整数约束',
               '求解：开源求解器 HiGHS；6 层场景按天数加权'], C['navy'], C['light'], size=16, minimum=15,
              head_h=0.42, name='right')
    T(s, x, 6.2, w, 0.7, ['方法来源：时刻表编制与车辆排班综述[41]；资源约束下的鲁棒时刻表[42]；线网与发车频率综述[43]；'
                          '多车型发车时刻与车型联合优化[44]；城乡一体化公交线路与频率联合优化[45]'], 12, 11, 'sources',
      color=C['grey'])


def p15(s, sp):
    lw = 6.05
    T(s, 0.5, 1.55, lw, 1.3,
      ['半柔性公交保留固定骨架、增加有限柔性[19]；柔性服务是公共交通的新机会[46]，但许多需求响应项目因成本过高或客流不足而停办[47]，'
       '因此柔性只用于不规则节律的日子'], 15, 14, 'idea')
    chain(s, ['次日预测（RQ2）', '前一晚指派', '次日运营'], 0.5, 2.95, lw, 0.55,
          colors=[C['teal'], C['orange'], C['navy']], size=14)
    # 4 车 × 2 半日格子
    gx, gy, c0, cw, rh = 0.5, 3.65, 2.0, (lw - 2.0 - 0.1) / 2, 0.42
    for j, t in enumerate(['上午 6:00—12:30', '下午 12:30—18:30']):
        lbox(s, gx + c0 + 0.05 + j * (cw + 0.05), gy, cw, rh, [t], fill=C['navy'], color=W, size=13, bold=True,
             name=f'grid_h{j + 1}')
    for i in range(4):
        y = gy + (i + 1) * (rh + 0.05)
        lbox(s, gx, y, c0, rh, [f'机动车 {i + 1}（19 座）'], fill=C['orange_l'], color=C['orange'], size=13, bold=True,
             name=f'grid_r{i + 1}')
        for j in range(2):
            lbox(s, gx + c0 + 0.05 + j * (cw + 0.05), y, cw, rh, ['线路 · 首发 · 往返'], fill=C['white'],
                 line=C['silver'], color=C['grey'], size=12, name=f'grid_{i + 1}{j + 1}')
    T(s, 0.5, 6.1, lw, 0.8, ['始末点均为县汽车站；求解按“车辆—半日”逐次贪心，每步选预测社会成本降幅最大且为正的候选'],
      14, 13, 'solve', color=C['navy'])
    # 右：对照与变体 + 信息价值
    x, w = 6.78, 6.05
    T(s, x, 1.55, w, 0.4, ['对照与变体（8 类）'], 16, 16, 'right_head', color=C['navy'], bold=True)
    alt = ['① 现行时刻表', '② 现行 + 调度员临时加班', '③ 只优化固定层', '④ 主方案（主方法预测）',
           '⑤ 改进预测方案', '⑥ 随机规划方案（样本平均近似，偏离预注册）[48]', '⑦ 反应式现场加开', '⑧ 完美预知上界']
    aw, ah = (w - 0.1) / 2, 0.66
    for i, t in enumerate(alt):
        ax = x + (i % 2) * (aw + 0.1); ay = 2.0 + (i // 2) * (ah + 0.05)
        hl = i in (3, 7)
        lbox(s, ax, ay, aw, ah, [t], fill=C['orange_l'] if hl else C['light'], line=C['orange'] if hl else None,
             color=C['text'], size=14, minimum=12, align='l', inset=0.06, bold=hl, name=f'alt{i + 1}')
    T(s, x, 4.8, w, 0.33, ['低需求条件下半柔性交付模式的评估见[49]'], 12, 12, 'alt_note', color=C['grey'])
    head_card(s, x, 5.15, w, 1.75, '信息价值：预测“兑现”了多少理论收益[21]',
              ['（主方案降幅 − 固定层降幅）÷（完美预知降幅 − 固定层降幅）'], C['orange'], C['orange_l'], size=17,
              minimum=15, head_h=0.42, bullet=False, name='formula', anchor='m')


def p16(s, sp):
    fo = sp['figure_options'][0]
    right, _ = fig(s, fo['full'], 0.5, 1.5, 7.45, 5.45, caption_of(sp, fo))
    x = right + 0.3; w = 12.83 - x
    box(s, x, 1.5, w, 0.38, fill=C['navy'], text='固定层（RQ3）', size=15, color=W, bold=True, align='l')
    T(s, x, 1.92, w, 1.75,
      ['检验期车次 36803 → 28648 班（减少 22.2%），车公里减少 24.9%', '服务率 81.6% → 83.8%，社会成本降低 13.8%',
       '代价：人均等待增加 2.5 分钟，满载班次率 3.0% → 4.8%'], 16, 15, 'fixed', bullet=True,
      emph=['22.2%', '13.8%'])
    box(s, x, 3.75, w, 0.38, fill=C['orange'], text='柔性层（RQ4）', size=15, color=W, bold=True, align='l')
    T(s, x, 4.17, w, 2.1,
      ['主方案社会成本降低 14.3%，比只优化固定层多 0.5 个百分点',
       '增量：节假日边沿 3.4、春运核心期 2.1、普通日 -0.2 个百分点',
       '信息价值：主方案 29.0%，改进预测 35.1%；完美预知上界降幅 15.4%',
       '建议：正式方案把“服务率不低于现行”作为硬约束'], 16, 15, 'flex', bullet=True,
      emph=['14.3%', '3.4', '2.1', '29.0%'])
    hyp_box(s, sp, x, 6.33, w, 0.6)


def p17(s, sp):
    cards = [('创新一 · 研究对象', C['teal'], C['teal_l'],
              '在“多重节律叠加”框架下，于县域尺度同时估计八个因子族的权重，为“集日是否重要”等地方经验提供可检验的定量证据'),
             ('创新二 · 方法', C['navy'], C['light'],
              '引入收益管理的删失需求还原[16]，构建“删失负二项—因子族 Shapley 分解—按情景 Mondrian 共形”链条，并在已知真值的模拟系统中检验其相对熵权法、灰色关联和 SHAP 的优势'),
             ('创新三 · 应用', C['orange'], C['orange_l'],
              '“固定层承担稳定节律、次日柔性层承担不规则节律”的半柔性配置[19]，用信息价值量化预测改进的边际收益，形成县级可操作的排班流程')]
    cw, gap = 3.98, (12.33 - 3 * 3.98) / 2
    for i, (h, col, bg, t) in enumerate(cards):
        head_card(s, 0.5 + i * (cw + gap), 1.6, cw, 2.7, h, [t], col, bg, size=15, minimum=14, head_h=0.48,
                  bullet=False, name=f'card{i + 1}')
    T(s, 0.5, 4.42, 12.33, 0.38, ['成果路线图'], 16, 16, 'road_head', color=C['navy'], bold=True)
    y, h = 4.85, 2.05
    pw, cwm = 3.85, 3.6
    cx = 0.5 + (12.33 - cwm) / 2
    for x, col, bg, head, body, nm in [
            (0.5, C['teal'], C['teal_l'], '论文一（RQ1—RQ2）',
             ['删失感知的多重节律识别与班次级概率预测', '面向交通需求建模与预测方向的国际期刊'], 'paper1'),
            (12.83 - pw, C['orange'], C['orange_l'], '论文二（RQ3—RQ4）',
             ['预测驱动的县域城乡公交半柔性配置', '面向公共交通运营与规划方向的国际期刊'], 'paper2')]:
        box(s, x, y, pw, h, fill=bg, line=col, line_w=1.25)
        T(s, x + 0.1, y + 0.08, pw - 0.2, 0.42, [head], 16, 16, nm + '_head', color=col, bold=True)
        T(s, x + 0.1, y + 0.55, pw - 0.2, h - 0.62, body, 15, 14, nm + '_body', bullet=True)
    box(s, cx, y, cwm, h, fill=C['navy'])
    T(s, cx + 0.1, y + 0.08, cwm - 0.2, 0.45, ['学位论文'], 18, 18, 'thesis_head', color=W, bold=True, align='c')
    T(s, cx + 0.1, y + 0.6, cwm - 0.2, h - 0.68, ['四个研究问题合成一篇完整的中文学位论文', '样稿已形成（见 04_综合研究报告）'],
      15, 14, 'thesis_body', color=W, align='c')
    arrow(s, 0.5 + pw + 0.05, y + h / 2, cx - 0.05, y + h / 2, color=C['teal'], width=2.25)
    arrow(s, 12.83 - pw - 0.05, y + h / 2, cx + cwm + 0.05, y + h / 2, color=C['orange'], width=2.25)


def p18(s, sp):
    cols = [('已完成', C['green'], W, ['文献综述（206 条参考文献，含中文核心期刊）', '研究问题、假设与预注册', '模拟系统、四个模型的全部代码与回测',
                                        '四份研究文档（研究方案、文献综述、计算模型研究报告、学位论文样稿）', '暑期跟车调查的现场工作']),
            ('进行中', C['yellow'], C['navy'], ['暑期跟车原始问卷录入', '与运营单位对接逐班次台账', '［请汇报人按实际情况确认］']),
            ('待开展', C['silver'], C['navy'], ['寒假与春运跟车调查', '行政村入户调查', '线上问卷', '用真实数据替换模拟数据并重算全部结果'])]
    tw = (7.4 - 2 * 0.08) / 3
    for i, (h, col, hc, items) in enumerate(cols):
        head_card(s, 0.5 + i * (tw + 0.08), 1.6, tw, 4.0, h, items, col, C['light'], size=14, minimum=12,
                  head_h=0.45, head_color=hc, name=f'prog{i + 1}')
    head_card(s, 8.15, 1.6, 4.68, 1.85, '待确认事项',
              ['各乡镇集期；天气阈值', '线路运行时间与配车、各车型核定座位数（目前为假设值）'], C['navy'], C['light'],
              size=14, minimum=13, head_h=0.42, head_size=15, name='todo')
    head_card(s, 8.15, 3.55, 4.68, 2.05, '研究局限',
              ['结果全部来自模拟数据', '只有全程总登乘人数，无站点上下客和 OD', '训练期只含两个春运',
               '两处偏离预注册的方案未计算预注册版本；集期与天气阈值待核实'], C['grey'], C['light'],
              size=12, minimum=11, head_h=0.42, head_size=15, name='limits')
    # 六阶段时间轴
    stages = [('第 1—2 月', '文献综述与数据对接', 1), ('第 3 月', '暑期跟车问卷与集期核实', 1),
              ('第 4—5 月', 'RQ1、RQ2 与论文一初稿', 0), ('第 6—7 月', '寒假（春运）跟车、入户与网络问卷', 0),
              ('第 8—9 月', 'RQ3、RQ4 与论文二初稿', 0), ('第 10—12 月', '学位论文撰写、修改与答辩', 0)]
    segs = chain(s, [''] * 6, 0.5, 6.05, 12.33, 0.85,
                 colors=[C['yellow'] if st else C['silver'] for _, _, st in stages])
    for (x, w), (m, t, st) in zip(segs, stages):
        T(s, x + 0.05, 5.72, w - 0.1, 0.3, [m], 11, 11, 'tl_month', color=C['navy'], bold=True, align='c')
        T(s, x + 0.3, 6.05, w - 0.5, 0.85, [t], 12, 11, 'timeline', color=C['navy'], bold=True, align='c', anchor='m')


def p_refs(s, sp):
    """参考文献两栏：著录中的方括号（[1]、[J]、[2026-09-30]）不做上标，编号悬挂缩进。"""
    from pptx.util import Pt
    b = sp['bullets']; h = [est_height_in([x], 9, 6.05, line=1.0, space_after=0.2) for x in b]
    cut = min(range(1, len(b)), key=lambda k: abs(sum(h[:k]) - sum(h[k:])))
    keep, pptkit.REF_RE = pptkit.REF_RE, re.compile(r'(?!)')
    REMAP[0] = False
    try:
        for (x, nm), part in (((0.5, 'col1'), b[:cut]), ((6.78, 'col2'), b[cut:])):
            tb, size, _ = T(s, x, 1.55, 6.05, 5.4, part, SIZE['ref'], SIZE['ref_min'], nm, line=1.0, space_after=0.2)
            for para in tb.text_frame.paragraphs:
                pPr = para._p.get_or_add_pPr(); ind = int(Pt(size * 2.4))
                pPr.set('marL', str(ind)); pPr.set('indent', str(-ind))
    finally:
        pptkit.REF_RE = keep
        REMAP[0] = True



# ===== 第四轮改写的页面（由 ppt/patch_build_r4.py 写入 build_ppt.py；文字中的文献编号为第一轮编号，输出时统一换成新编号） =====
def split_label(t, sep='：'):
    k = t.index(sep)
    return t[:k], t[k + 1:]


def plain_banner(s, text, y=1.55, h=0.6, name='plain'):
    """“通俗地说”通栏：浅青底 + 左侧色条，开头四个字着色。"""
    box(s, 0.5, y, 12.33, h, fill=C['teal_l'])
    box(s, 0.5, y, 0.07, h, fill=C['teal'])
    lab, txt = split_label(text)
    T(s, 0.65, y, 12.1, h, [lab + '：' + txt], 15, 13, name, color=C['navy'], anchor='m', emph=[lab + '：'])


def p3(s, sp):
    items = [('政策要求', C['navy'],
              '2011 年起推进城乡道路客运一体化[1]；2024 年服务指南要求按客流规律灵活安排班次[2]；2026 年“十五五”枢纽规划要求运行时刻对接、客流高峰期预警和适老化[50]'),
             ('山区现实', C['teal'],
              '交通劣势与相对贫困存在可测的关联[3]；医疗等公共服务的可达性明显受交通条件限制[4]'),
             ('理论视角', C['grey'],
              '缺少公交使农村居民难以到达就业、医疗和教育机会，属于社会排斥问题[5,6]；城乡交通融合已成为地理学研究议题[7]')]
    hs = [2.0, 1.43, 1.63]
    y = 1.6
    for i, ((lab, col, txt), bh) in enumerate(zip(items, hs)):
        box(s, 0.5, y, 6.05, bh, fill=C['light'])
        box(s, 0.5, y, 1.3, bh, fill=col, text=lab, size=17, color=W, bold=True)
        T(s, 1.9, y + 0.04, 4.6, bh - 0.08, [txt], 16, 15, f'left{i + 1}', anchor='m')
        y += bh + 0.12
    x, y, w, h = 6.78, 1.6, 6.05, 5.3
    box(s, x, y, w, h, fill=C['light'])
    box(s, x, y, w, 0.45, fill=C['navy'], text='研究对象', size=16, color=W, bold=True, align='l')
    T(s, x + 0.15, y + 0.52, w - 0.3, 0.56, ['湖南省新晃侗族自治县（武陵山区）'], 20, 18, 'obj1', color=C['navy'], bold=True)
    T(s, x + 0.15, y + 1.08, w - 0.3, 0.42, ['县汽车站发出的 7 条城乡公交线路'], 16, 16, 'obj2', color=C['teal'], bold=True)
    fo = SP[8]['figure_options'][0]
    fig(s, fo['full'], x + 0.15, y + 1.55, w - 0.3, h - 1.6, '图：研究区与城乡公交线网（真实 GIS 底图与现行线路，详见第 8 页）',
        halign='c')


def p6(s, sp):
    mx, my, mw, mh = 0.5, 1.6, 7.7, 4.05
    rq = []
    for t in sp['bullets'][:4]:
        k = t.index('？')
        rq.append((t[:3] + '　' + t[4:k + 1], t[k + 1:]))
    lw, gap = 0.55, 0.1
    cw, ch = (mw - lw - 2 * gap) / 2, (mh - gap) / 2
    for r, (lab, col, bg) in enumerate([('需求', C['teal'], C['teal_l']), ('供给', C['orange'], C['orange_l'])]):
        y = my + r * (ch + gap)
        box(s, mx, y, lw, ch, fill=col, text=list(lab), size=18, color=W, bold=True)
        for c in range(2):
            k = r * 2 + c; x = mx + lw + gap + c * (cw + gap)
            head_card(s, x, y, cw, ch, rq[k][0], [rq[k][1]], col, bg, size=15, minimum=13, bullet=False,
                      name=f'rq{k + 1}')
    hx, hw, gy = 8.45, 4.38, 0.07
    T(s, hx, 1.55, hw, 0.4, ['我们的预期（假设，可被数据否定）'], 15, 14, 'hyp_head', color=C['navy'], bold=True)
    hyp = [('H1', 'RQ1', '“几点钟”影响最大；春运、节假日决定哪几天特别多；赶集日、天气影响较小'),
           ('H2', 'RQ1', '常用的熵权法、灰色关联给出的排序与真实情况明显不一致'),
           ('H3', 'RQ2', '把“坐满”考虑进去的模型，对满员班次低估更少'),
           ('H4', 'RQ2', '按春运、节假日、平日分别给出的人数范围，比“一刀切”的更可靠'),
           ('H5', 'RQ3', '只按日常规律编排的固定时刻表，班次更少、总成本更低'),
           ('H6', 'RQ4', '机动车的好处集中在春运和节假日前后，平常日子几乎没有')]
    y0 = 2.0
    hh = (6.9 - y0 - 5 * gy) / 6
    for i, (h, r, t) in enumerate(hyp):
        y = y0 + i * (hh + gy); col = C['teal'] if i < 4 else C['orange']
        box(s, hx, y, hw, hh, fill=C['teal_l'] if i < 4 else C['orange_l'])
        box(s, hx, y, 0.62, hh, fill=col, text=[h, r], size=12, color=W, bold=True)
        T(s, hx + 0.66, y, hw - 0.7, hh, [t], 14, 12, f'hyp{i + 1}', anchor='m')
    T(s, mx, 5.7, mw, 0.3, ['两篇论文合起来构成学位论文'], 12, 12, 'papers_head', color=C['navy'], bold=True, align='c')
    box(s, mx, 6.0, mw, 0.03, fill=C['navy'])
    pw = (mw - 0.1) / 2
    lbox(s, mx, 6.08, pw, 0.82, ['论文一（RQ1—RQ2）', '删失感知的多重节律识别与概率预测'], fill=C['teal'], color=W,
         size=13, bold=True, name='paper1')
    lbox(s, mx + pw + 0.1, 6.08, pw, 0.82, ['论文二（RQ3—RQ4）', '预测驱动的县域城乡公交半柔性配置'], fill=C['orange'],
         color=W, size=13, bold=True, name='paper2')


def p7(s, sp):
    b = sp['bullets']
    nodes = []
    for t in b[:6]:
        head, rest = split_label(t)
        body, _, out = rest.partition('得到：')
        nodes.append((head, body.strip(), out.strip().rstrip('。')))
    cols = [C['grey'], C['teal'], C['teal'], C['orange'], C['orange'], C['navy']]
    bgs = [C['light'], C['teal_l'], C['teal_l'], C['orange_l'], C['orange_l'], C['light']]
    g = 0.25; nw = (12.33 - 5 * g) / 6; y0, nh = 1.6, 2.75
    for i, (head, body, out) in enumerate(nodes):
        x = 0.5 + i * (nw + g)
        box(s, x, y0, nw, nh, fill=bgs[i])
        lbox(s, x, y0, nw, 0.62, [head], fill=cols[i], color=W, size=14, minimum=12, bold=True, name=f'node{i + 1}_h',
             inset=0.04)
        oh = 0.62 if out else 0
        T(s, x + 0.06, y0 + 0.68, nw - 0.12, nh - 0.74 - oh - 0.06, [body.rstrip('。')], 13, 11, f'node{i + 1}')
        if out:
            lbox(s, x + 0.06, y0 + nh - oh - 0.06, nw - 0.12, oh, ['得到：' + out], fill=W, line=cols[i],
                 color=cols[i], size=12, minimum=10.5, bold=True, name=f'node{i + 1}_out', inset=0.03)
        if i < 5:
            arrow(s, x + nw + 0.02, y0 + nh / 2, x + nw + g - 0.02, y0 + nh / 2, color=C['navy'], width=2.25)
    # 需求侧 / 供给侧括注
    yb = y0 + nh + 0.1
    for (i1, i2, lab, col) in ((1, 2, '需求是什么 → 论文一', C['teal']), (3, 4, '供给怎么配 → 论文二', C['orange'])):
        xa = 0.5 + i1 * (nw + g); xb = 0.5 + i2 * (nw + g) + nw
        box(s, xa, yb, xb - xa, 0.03, fill=col)
        T(s, xa, yb + 0.05, xb - xa, 0.34, [lab], 13, 12, f'brace{i1}', color=col, bold=True, align='c')
    # 三条规则
    rules = [split_label(t) for t in b[6:9]]
    rw = (12.33 - 2 * 0.2) / 3; ry = 4.92
    for i, (h, t) in enumerate(rules):
        head_card(s, 0.5 + i * (rw + 0.2), ry, rw, 6.92 - ry, h, [t], C['navy'], C['light'], size=14, minimum=12,
                  head_h=0.42, head_size=15, name=f'rule{i + 1}', bullet=False)


def p9(s, sp):
    b = sp['bullets']
    plain_banner(s, b[0], y=1.48, h=0.74, name='why')
    steps = [split_label(t) for t in b[1:5]]
    g = 0.28; sw = (12.33 - 3 * g) / 4; sy, sh = 2.3, 1.8
    for i, (h, t) in enumerate(steps):
        x = 0.5 + i * (sw + g)
        head_card(s, x, sy, sw, sh, h, [t], C['navy'] if i != 3 else C['teal'], C['light'], size=13, minimum=11,
                  head_h=0.38, head_size=14, name=f'step{i + 1}', bullet=False)
        if i < 3:
            arrow(s, x + sw + 0.03, sy + sh / 2, x + sw + g - 0.03, sy + sh / 2, color=C['navy'], width=2.25)
    ly = sy + sh + 0.14
    right, bottom = fig(s, sp['figures'][0], 0.5, ly, 5.75, 6.95 - ly - 0.72, '图：模拟数据核对，面板 a（模拟数据，SYNTHETIC）')
    lbox(s, 0.5, bottom + 0.06, 5.75, 6.95 - bottom - 0.06, [b[7]], fill=C['teal_l'], color=C['navy'], size=13,
         minimum=11, bold=True, align='l', inset=0.08, name='real_data')
    x = 0.5 + 5.75 + 0.3; w = 12.83 - x
    cards = [('852', '天', '模拟天数'), ('116127', '个', '班次观测'), ('3.0', '%', '满载删失班次'), ('255', '次', '调度员临时加班'),
             ('808', '份', '模拟问卷')]
    cw = (w - 0.08 * 4) / 5; chh = 0.82
    for i, (v, u, lab) in enumerate(cards):
        bx = x + i * (cw + 0.08)
        box(s, bx, ly, cw, chh, fill=C['light'])
        T(s, bx, ly + 0.02, cw, 0.46, [v + u], 18, 13, 'num' + str(i + 1), color=C['orange'], bold=True, align='c',
          anchor='m')
        T(s, bx, ly + 0.48, cw, 0.32, [lab], 10.5, 10, 'numlab' + str(i + 1), color=C['grey'], align='c')
    head, txt = split_label(b[6])
    items = [x.strip().rstrip('。') for x in re.split(r'；(?=RQ\d)', txt)]
    head_card(s, x, ly + chh + 0.12, w, 6.95 - ly - chh - 0.12, head + '（模拟数据）', items, C['orange'], C['orange_l'],
              size=13, minimum=10.5, head_h=0.38, head_size=14, name='summary',
              )


def model_plain(s, sp, steps, why_head, why, side_head, side, side_style='plain', step_cols=None):
    """模型页通用版式：通俗通栏 + 三步流程 + 下方“为什么/怎么比”与侧栏。"""
    plain_banner(s, sp['bullets'][0])
    g = 0.45; sw = (12.33 - 2 * g) / 3; sy, sh = 2.3, 1.9
    step_cols = step_cols or [C['grey'], C['teal'], C['navy']]
    for i, (h, t) in enumerate(steps):
        x = 0.5 + i * (sw + g)
        head_card(s, x, sy, sw, sh, h, [t], step_cols[i], C['light'], size=14, minimum=12, head_h=0.42, head_size=15,
                  name=f'step{i + 1}', bullet=False)
        if i < 2:
            arrow(s, x + sw + 0.05, sy + sh / 2, x + sw + g - 0.05, sy + sh / 2, color=C['navy'], width=2.5)
    ly = sy + sh + 0.15; lw = 7.65
    head_card(s, 0.5, ly, lw, 6.95 - ly, why_head, why, C['teal'], C['teal_l'], size=14, minimum=12, head_h=0.42,
              head_size=15, name='why')
    x = 0.5 + lw + 0.2; w = 12.83 - x
    if side_style == 'dashed':
        box(s, x, ly, w, 6.95 - ly, fill=W, line=C['orange'], dash=True, line_w=1.5)
        T(s, x + 0.1, ly + 0.06, w - 0.2, 0.4, [side_head], 15, 14, 'side_head', color=C['orange'], bold=True)
        T(s, x + 0.1, ly + 0.5, w - 0.2, 6.95 - ly - 0.58, side, 14, 12, 'side')
    else:
        head_card(s, x, ly, w, 6.95 - ly, side_head, side, C['grey'], C['light'], size=13, minimum=11, head_h=0.42,
                  head_size=14, name='side', bullet=False)


def p10(s, sp):
    b = sp['bullets']
    steps = [('输入', split_label(b[1])[1]), ('怎么算', split_label(b[2])[1]), ('得到', split_label(b[3])[1])]
    why = ['客流是 0、1、2……这样的整数，而且忽多忽少，负二项模型专门描述这种数据[28]',
           '坐满的班次若当成“正好 c 人”会低估节点日需求，航空收益管理早有“需求还原”的做法[29]',
           '“公平分配”用 Shapley 方法：各因素份额加起来正好等于总数，因素相关也不重复计算[30,31]',
           '和常用方法比：熵权法、灰色关联只看波动大小或形状是否相似，不衡量对客流的解释力；SHAP 在因素相关时分配会失真[32]']
    model_plain(s, sp, steps, '为什么这样做（和常用方法比）', why, '专业名称（供提问参考）', [split_label(b[7])[1]])


def p12(s, sp):
    b = sp['bullets']
    steps = [split_label(t) for t in b[1:4]]
    why = [b[4], b[5], b[6]]
    lab, txt = split_label(b[7])
    model_plain(s, sp, steps, '为什么分情景、怎么检验', why, lab, [txt], side_style='dashed',
                step_cols=[C['teal'], C['teal'], C['navy']])


def p14(s, sp):
    b = sp['bullets']
    plain_banner(s, b[0])
    lw = 6.05; y0 = 2.3
    lbox(s, 0.5, y0, (lw - 0.1) / 2, 0.85, ['固定时刻表：管“天天都有”的规律——钟点、星期几、上学或放假'], fill=C['navy'],
         color=W, size=13, minimum=11, bold=True, name='split_fixed', inset=0.06)
    lbox(s, 0.5 + (lw + 0.1) / 2, y0, (lw - 0.1) / 2, 0.85, ['次日机动车：管“偶尔才有”的高峰——春运、节假日、恶劣天气'],
         fill=C['orange_l'], line=C['orange'], color=C['orange'], size=13, minimum=11, bold=True, name='split_flex',
         inset=0.06)
    T(s, 0.5, y0 + 0.92, lw, 0.34, ['六类日子一起排，按天数加权'], 14, 13, 'layers_head', color=C['navy'], bold=True)
    layers = ['学期周一至周四', '学期周五', '学期周六', '学期周日', '假期工作日', '假期周末']
    cw = (lw - 2 * 0.08) / 3
    for i, t in enumerate(layers):
        x = 0.5 + (i % 3) * (cw + 0.08); y = y0 + 1.3 + (i // 3) * 0.48
        lbox(s, x, y, cw, 0.42, [t], fill=C['teal_l'] if i < 4 else C['light'], line=C['teal'] if i < 4 else C['silver'],
             color=C['navy'], size=13, name=f'layer{i + 1}')
    sw = (lw - 0.4) / 2; iy = y0 + 2.32
    lbox(s, 0.5, iy, sw, 0.8, ['估计每班真实需求', '（沿用模型一）'], fill=C['teal'], color=W, size=13, bold=True, name='inv1')
    arrow(s, 0.5 + sw + 0.05, iy + 0.4, 0.5 + sw + 0.35, iy + 0.4, color=C['navy'], width=2)
    lbox(s, 0.5 + sw + 0.4, iy, sw, 0.8, ['按问卷“想几点出发”', '分到每 15 分钟'], fill=C['teal'], color=W, size=13,
         bold=True, name='inv2')
    import math
    cy = iy + 0.9
    box(s, 0.5, cy, lw, 6.95 - cy, fill=C['light'])
    bx, by, bw, bh = 0.7, cy + 0.1, lw - 0.4, 6.95 - cy - 0.5
    n = 48
    for k in range(n):
        t = k / (n - 1)
        v = 0.15 + 0.85 * math.exp(-((t - 0.12) / 0.07) ** 2) + 0.55 * math.exp(-((t - 0.62) / 0.12) ** 2)
        hh = bh * min(v, 1.0)
        box(s, bx + k * bw / n, by + bh - hh, bw / n * 0.8, hh, fill=C['teal'] if v > 0.6 else C['silver'])
    box(s, bx, by + bh, bw, 0.015, fill=C['grey'])
    textbox(s, 0.6, 6.95 - 0.36, lw - 0.2, 0.3, ['15 分钟需求曲线（示意）：固定发车按此曲线在网格上排布'], size=11,
            color=C['grey'], fit=False)
    x, w = 6.78, 6.05
    _, goal = split_label(b[3])
    head_card(s, x, y0, w, 1.55, '目标：社会成本最低[39,40]', [goal.split('。', 1)[1].replace('[39,40]', '').rstrip('。')],
              C['orange'], C['orange_l'], size=15, minimum=13, head_h=0.42, bullet=False, name='formula', anchor='m')
    head_card(s, x, y0 + 1.65, w, 2.25, '决定什么、要满足什么',
              [b[4].rstrip('。'), b[5].rstrip('。')], C['navy'], C['light'], size=14, minimum=12, head_h=0.42,
              head_size=15, name='right')
    T(s, x, y0 + 3.98, w, 6.97 - y0 - 3.98, [b[6]], 12, 10, 'sources', color=C['grey'])


def p15(s, sp):
    b = sp['bullets']
    plain_banner(s, b[0])
    lw = 6.05
    T(s, 0.5, 2.25, lw, 1.15, [b[1]], 14, 12, 'idea')
    chain(s, ['次日预测（模型二）', '前一晚派车', '次日运营'], 0.5, 3.48, lw, 0.5,
          colors=[C['teal'], C['orange'], C['navy']], size=13)
    gx, gy, c0, cw, rh = 0.5, 4.08, 2.0, (lw - 2.0 - 0.1) / 2, 0.36
    for j, t in enumerate(['上午 6:00—12:30', '下午 12:30—18:30']):
        lbox(s, gx + c0 + 0.05 + j * (cw + 0.05), gy, cw, rh, [t], fill=C['navy'], color=W, size=12, bold=True,
             name=f'grid_h{j + 1}')
    for i in range(4):
        y = gy + (i + 1) * (rh + 0.04)
        lbox(s, gx, y, c0, rh, [f'机动车 {i + 1}（19 座）'], fill=C['orange_l'], color=C['orange'], size=12, bold=True,
             name=f'grid_r{i + 1}')
        for j in range(2):
            lbox(s, gx + c0 + 0.05 + j * (cw + 0.05), y, cw, rh, ['线路 · 首发 · 往返'], fill=C['white'],
                 line=C['silver'], color=C['grey'], size=11, name=f'grid_{i + 1}{j + 1}')
    sy = gy + 5 * (rh + 0.04) + 0.06
    _, how = split_label(b[2])
    T(s, 0.5, sy, lw, 6.95 - sy, ['怎么派：' + how.split('；', 1)[1]], 13, 11, 'solve', color=C['navy'])
    x, w = 6.78, 6.05
    T(s, x, 2.22, w, 0.38, ['和 8 种方案比较'], 15, 15, 'right_head', color=C['navy'], bold=True)
    _, alts = split_label(b[3])
    alts, _, note = alts.partition('。')
    alt = alts.split('；')
    aw, ah = (w - 0.1) / 2, 0.6
    for i, t in enumerate(alt):
        ax = x + (i % 2) * (aw + 0.1); ay = 2.64 + (i // 2) * (ah + 0.05)
        hl = i in (3, 7)
        lbox(s, ax, ay, aw, ah, [f'{"①②③④⑤⑥⑦⑧"[i]} {t}'], fill=C['orange_l'] if hl else C['light'],
             line=C['orange'] if hl else None, color=C['text'], size=13, minimum=10.5, align='l', inset=0.06, bold=hl,
             name=f'alt{i + 1}')
    T(s, x, 5.21, w, 0.32, [note.strip().rstrip('。')], 11.5, 11, 'alt_note', color=C['grey'])
    lab, formula = b[4].split('＝', 1)
    head_card(s, x, 5.55, w, 1.4, lab, [formula.rstrip('。')], C['orange'], C['orange_l'], size=15, minimum=13,
              head_h=0.42, bullet=False, name='formula', anchor='m')


def p17(s, sp):
    b = sp['bullets']
    cards = [(C['teal'], C['teal_l']), (C['navy'], C['light']), (C['orange'], C['orange_l'])]
    cw, gap = 3.98, (12.33 - 3 * 3.98) / 2
    for i, ((col, bg), t) in enumerate(zip(cards, b[:3])):
        h, body = split_label(t)
        head_card(s, 0.5 + i * (cw + gap), 1.6, cw, 2.2, h.replace('（', ' · ').rstrip('）'), [body.rstrip('。')], col,
                  bg, size=14, minimum=12, head_h=0.45, bullet=False, name=f'card{i + 1}')
    head, rest = split_label(b[3])
    T(s, 0.5, 3.92, 12.33, 0.4, [head], 16, 15, 'plan_head', color=C['navy'], bold=True)
    pts = [p for p in re.split(r'[；]?[①②③④]', rest) if p.strip()]
    bw = (12.33 - 3 * 0.15) / 4; by = 4.38
    for i, p in enumerate(pts):
        req, out = [x.strip().rstrip('。') for x in p.split('→')]
        x = 0.5 + i * (bw + 0.15)
        lbox(s, x, by, bw, 0.8, [f'{"①②③④"[i]} {req}'], fill=C['navy'], color=W, size=13, minimum=11, bold=True,
             name=f'plan{i + 1}_req', inset=0.06)
        arrow(s, x + bw / 2, by + 0.82, x + bw / 2, by + 1.06, color=C['navy'], width=2)
        col, bg = (C['teal'], C['teal_l']) if i < 3 else (C['orange'], C['orange_l'])
        lbox(s, x, by + 1.08, bw, 0.72, [out], fill=bg, line=col, color=col, size=13, minimum=11, bold=True,
             name=f'plan{i + 1}_out', inset=0.06)
    lbox(s, 0.5, 6.3, 12.33, 0.65, [b[4]], fill='FFF4D1', line=C['yellow'], color=C['text'], size=13, minimum=11,
         align='l', inset=0.1, name='plan_pos')


def p18(s, sp):
    cols = [('已完成', C['green'], W, ['文献综述（206 条参考文献，含中文核心期刊）', '研究问题、假设与预注册', '模拟系统、四个模型的全部代码与回测',
                                        '四份研究文档（研究方案、文献综述、计算模型研究报告、学位论文样稿）', '暑期跟车调查的现场工作']),
            ('进行中', C['yellow'], C['navy'], ['暑期跟车原始问卷录入', '与运营单位对接逐班次台账', '［请汇报人按实际情况确认］']),
            ('待开展', C['silver'], C['navy'], ['寒假与春运跟车调查', '行政村入户调查', '线上问卷', '用真实数据替换模拟数据并重算全部结果'])]
    tw = (7.4 - 2 * 0.08) / 3
    for i, (h, col, hc, items) in enumerate(cols):
        head_card(s, 0.5 + i * (tw + 0.08), 1.6, tw, 4.0, h, items, col, C['light'], size=14, minimum=12,
                  head_h=0.45, head_color=hc, name=f'prog{i + 1}')
    head_card(s, 8.15, 1.6, 4.68, 2.2, '待确认事项',
              ['各乡镇集期；天气阈值', '线路运行时间与配车、各车型核定座位数（目前为假设值）',
               '是否纳入“十五五”规划相关的补强方向：与铁路客站到发时刻衔接、客货邮融合、能耗测算'], C['navy'], C['light'],
              size=13, minimum=11, head_h=0.42, head_size=15, name='todo')
    head_card(s, 8.15, 3.88, 4.68, 1.72, '研究局限',
              ['结果全部来自模拟数据', '只有全程总登乘人数，无站点上下客和 OD', '训练期只含两个春运',
               '两处偏离预注册的方案未计算预注册版本；集期与天气阈值待核实'], C['grey'], C['light'],
              size=12, minimum=10, head_h=0.4, head_size=15, name='limits')
    stages = [('第 1—2 月', '文献综述与数据对接', 1), ('第 3 月', '暑期跟车问卷与集期核实', 1),
              ('第 4—5 月', 'RQ1、RQ2 与论文一初稿', 0), ('第 6—7 月', '寒假（春运）跟车、入户与网络问卷', 0),
              ('第 8—9 月', 'RQ3、RQ4 与论文二初稿', 0), ('第 10—12 月', '学位论文撰写、修改与答辩', 0)]
    segs = chain(s, [''] * 6, 0.5, 6.05, 12.33, 0.85,
                 colors=[C['yellow'] if st else C['silver'] for _, _, st in stages])
    for (x, w), (m, t, st) in zip(segs, stages):
        T(s, x + 0.05, 5.72, w - 0.1, 0.3, [m], 11, 11, 'tl_month', color=C['navy'], bold=True, align='c')
        T(s, x + 0.3, 6.05, w - 0.5, 0.85, [t], 12, 11, 'timeline', color=C['navy'], bold=True, align='c', anchor='m')



PAGES = {1: p1, 2: p2, 3: p3, 4: p4, 5: p5, 6: p6, 7: p7, 8: p8, 9: p9, 10: p10, 11: p11, 12: p12,
         13: p13, 14: p14, 15: p15, 16: p16, 17: p17, 18: p18, 19: p_refs, 20: p_refs}

R1 = {sl['no']: sl for sl in json.load(open('ppt/ppt_spec_r1.json', encoding='utf-8'))['slides']}
NOTE_EXTRA = {
    3: '右栏插图与第 8 页为同一张真实 GIS 底图，此处用于交代研究对象。',
    7: '技术路线改用原生形状绘制的六步流程；完整技术路线图见 fig/图_技术路线.png 与研究方案文档。',
    16: '甘特图样例（fig/图_次日柔性排班甘特图样例.png，宽高比 3.0）未放入页面，答问时可展示：它给出某一天 4 辆机动车上午、下午两个半日的线路指派与发车时刻。',
}

for _no in (6, 7, 10, 12, 14, 15):           # 改为通俗表述的页面：备注中保留第一轮的技术表述，供答问时使用
    _orig = '第一轮的技术表述（供答问参考）：\n' + '\n'.join(cite_remap(x) for x in R1[_no]['bullets'])
    NOTE_EXTRA[_no] = (NOTE_EXTRA[_no] + '\n\n' + _orig) if _no in NOTE_EXTRA else _orig

prs = new_prs()
for sp in spec['slides']:
    s = blank(prs); no = sp['no']; NO[0] = no
    if sp['layout_id'] != 'cover':
        # 标题框上下内边距取 0，使两行标题在 y = 0.30—1.36 英寸内仍能保持 ≥ 26 磅（第 4 节字号要求）
        vm, pptkit.vmargin = pptkit.vmargin, lambda h: 0.0
        SIZE['title_min'] = 26
        try:
            # 标题在 28 磅下只多出一两个字时，降到 26—27 磅放成一行，避免第二行孤字
            tw = (12.33 - 2 * pptkit.MARGIN) * 72
            ts = next((z for z in (28, 27.5, 27, 26.5, 26) if len(wrap_lines(sp['title'], z, tw, True)) == 1), 28)
            frame(s, no, TOTAL, sp['title'], sp['section'], footer=SYN if sp['synthetic'] else None, title_size=ts)
        finally:
            pptkit.vmargin = vm
    PAGES[no](s, sp if sp['layout_id'] == 'refs' else SP[no])      # 参考文献页用新编号著录
    extra = '\n\n'.join(x for x in (SYN if sp['synthetic'] else None, NOTE_EXTRA.get(no)) if x) or None
    set_notes(s, sp['speaker_notes'] or '', sp['bullets'] if no < 19 else None, extra)
prs.save(OUT)
json.dump(pptkit.FIT_LOG, open('ppt/fit_log.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
over = [f"{d['slide']}:{d['box']}" for d in pptkit.FIT_LOG if d['overflow']]
print(OUT, len(prs.slides), '页；溢出文本框', len(over), '个：', ' '.join(over))
