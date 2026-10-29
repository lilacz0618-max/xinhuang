"""第三轮 PPT 自检（self_check_v3.py）：与 self_check_v2.py 相同的检查项，
只把 spec 换成 ppt/ppt_spec_r3.json、页数上限改为 27（新增两页“十五五”规划对接），报告写入 ppt_check_r3.md。
原第二轮说明：
用法：python ppt/tools/self_check_v2.py 新晃城乡公交研究进展汇报_第二轮.pptx
依赖：ppt/ppt_spec_r2.json、tab/numbers.json；可选 ppt/fit_log.json、preview_quick/render_issues.json
输出：ppt_check_r2.md；“失败”必须为 0 才能交付。
"""
import sys, os, re, json, io
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image

PPT = sys.argv[1] if len(sys.argv) > 1 else '新晃城乡公交研究进展汇报_第三轮.pptx'
SPEC = 'ppt/ppt_spec_r3.json'
FIT = 'ppt/fit_log.json'
REN = 'preview_quick/render_issues.json'
SYN = 'SYNTHETIC'
MAX_PAGES = 27
BLOCKS = ['要回答的问题', '为什么选这个方法', '和其他方法比', '怎么证明它有效']

spec = json.load(open(SPEC, encoding='utf-8'))
nums = json.load(open('tab/numbers.json', encoding='utf-8'))
nref = len(spec['references'])
prs = Presentation(PPT)
fails, warns, rows = [], [], []


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def shape_text(sh):
    t = []
    if getattr(sh, 'has_text_frame', False) and sh.has_text_frame:
        t.append(sh.text_frame.text)
    if getattr(sh, 'has_table', False) and sh.has_table:
        for r in sh.table.rows:
            for c in r.cells:
                t.append(c.text)
    return '\n'.join(t)


def shape_sizes(sh):
    tfs = []
    if getattr(sh, 'has_text_frame', False) and sh.has_text_frame:
        tfs.append(sh.text_frame)
    if getattr(sh, 'has_table', False) and sh.has_table:
        tfs += [c.text_frame for r in sh.table.rows for c in r.cells]
    out = []
    for tf in tfs:
        for p in tf.paragraphs:
            for r in p.runs:
                if r.font.size is not None and r.text.strip():
                    out.append(r.font.size.pt)
    return out


def cites_in(text):
    return sorted({int(x) for g in re.findall(r'\[(\d+(?:[,，–-]\d+)*)\]', text)
                   for x in re.split(r'[,，–-]', g) if x})


n = len(prs.slides)
if n > MAX_PAGES:
    fails.append(f'总页数 {n} 超过 {MAX_PAGES}')
if n != len(spec['slides']):
    fails.append(f'PPT 页数 {n} 与 spec 页数 {len(spec["slides"])} 不一致')

seen_max = 0
for sl, sp in zip(prs.slides, spec['slides']):
    no, lay = sp['no'], sp['layout_id']
    shapes = list(iter_shapes(sl.shapes))
    foot = [s for s in shapes if s.name == 'footnote']
    plain = [s for s in shapes if s.name == 'plain']
    body = '\n'.join(shape_text(s) for s in shapes if s.name != 'footnote')
    foot_t = '\n'.join(shape_text(s) for s in foot)
    text = body + '\n' + foot_t
    notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ''
    # 1 数字
    for k, v in sp['numbers'].items():
        if str(nums.get(k)) != str(v):
            fails.append(f'第{no}页 数字键 {k} 与 tab/numbers.json 不一致')
        if str(v) not in body:
            fails.append(f'第{no}页 数字 {k}={v} 未出现在页面上')
    # 2 文献：正文引用、同页脚注、首次出现顺序
    if lay == 'refs':
        for i in sp.get('ref_range', []):
            if f'[{i}]' not in body:
                fails.append(f'第{no}页 参考文献总表缺少 [{i}]')
        bc = []
    else:
        bc = cites_in(body)
        if any(not 1 <= c <= nref for c in bc):
            fails.append(f'第{no}页 文献编号越界：{bc}')
        miss = sorted(set(sp['refs']) - set(bc))
        if miss:
            fails.append(f'第{no}页 spec 要求的文献 {miss} 没有在正文中标注')
        extra = sorted(set(bc) - set(sp['refs']))
        if extra:
            fails.append(f'第{no}页 正文出现 spec 之外的文献编号 {extra}')
        if sp['refs']:
            if not foot:
                fails.append(f'第{no}页 有引用但没有名为 footnote 的同页文献脚注框')
            else:
                fc = cites_in(foot_t)
                lack = sorted(set(sp['refs']) - set(fc))
                if lack:
                    fails.append(f'第{no}页 脚注缺少 {lack}')
            if '本页文献' not in notes:
                fails.append(f'第{no}页 备注缺少“本页文献”完整著录段')
        new = [c for c in sp['refs'] if c > seen_max]
        if new and min(new) != seen_max + 1:
            fails.append(f'第{no}页 文献未按首次出现顺序编号（应从 [{seen_max + 1}] 开始）')
        seen_max = max([seen_max] + sp['refs'])
    # 3 SYNTHETIC
    if sp['synthetic'] and SYN not in body:
        fails.append(f'第{no}页 为模拟结果页，页面缺少 SYNTHETIC 页脚')
    if sp['synthetic'] and SYN not in notes:
        fails.append(f'第{no}页 讲稿备注缺少 SYNTHETIC 说明')
    if no == 1 and SYN not in body:
        fails.append('封面缺少 SYNTHETIC 说明')
    # 4 术语与表述
    if '截断' in text + notes:
        fails.append(f'第{no}页 出现“截断”，应为“删失”')
    for para in re.split(r'[。\n；]', body):
        if ('集日' in para or '天气' in para) and re.search(r'核心(因|变量|依据)', para) \
                and not re.search(r'不(宜|应|作|是)|对照', para):
            warns.append(f'第{no}页 疑似把集日/天气写成核心因素：“{para.strip()[:40]}”')
    for s in sp.get('must_include', []):
        if s not in body:
            fails.append(f'第{no}页 缺少必须保留的表述：“{s}”')
    if sp.get('plain'):
        if not plain or '通俗地说' not in shape_text(plain[0]):
            fails.append(f'第{no}页 缺少名为 plain、以“通俗地说”开头的通俗解释框')
    if lay == 'model4':
        for b in BLOCKS:
            if b not in body:
                fails.append(f'第{no}页 模型页缺少“{b}”区块')
    if '偏离预注册' in json.dumps(sp, ensure_ascii=False) and '偏离预注册' not in body + notes:
        fails.append(f'第{no}页 删去了“偏离预注册”字样')
    # 5 图片
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                im = Image.open(io.BytesIO(sh.image.blob))
                r0, r1 = im.width / im.height, sh.width / sh.height
                if abs(r1 / r0 - 1) > 0.01:
                    fails.append(f'第{no}页 图片被拉伸（原比例 {r0:.3f}，现比例 {r1:.3f}）')
                if sh.crop_left or sh.crop_right or sh.crop_top or sh.crop_bottom:
                    fails.append(f'第{no}页 图片被裁剪')
                if sp.get('figures') and lay == 'result_fig' and sh.width / prs.slide_width < 0.28:
                    warns.append(f'第{no}页 结果图宽度不足页面 28%，可能看不清')
            except Exception as e:
                warns.append(f'第{no}页 无法读取图片：{e}')
    # 6 字号
    bs = [x for s in shapes if s.name != 'footnote' for x in shape_sizes(s)]
    fs = [x for s in foot for x in shape_sizes(s)]
    lim = 9 if lay == 'refs' else 10
    if bs and min(bs) < lim:
        fails.append(f'第{no}页 正文最小字号 {min(bs)} 磅低于 {lim} 磅')
    if fs and min(fs) < 8.5:
        fails.append(f'第{no}页 脚注字号 {min(fs)} 磅低于 8.5 磅')
    # 7 备注
    if lay != 'refs' and len(notes.strip()) < 30:
        fails.append(f'第{no}页 讲稿备注为空或过短')
    if lay not in ('refs', 'cover') and '完整要点：' not in notes:
        fails.append(f'第{no}页 备注缺少“完整要点：”段')
    # 8 占位
    if 'TODO' in text or '待替换' in text:
        fails.append(f'第{no}页 仍有 TODO/占位框')
    if '［' in body:
        warns.append(f'第{no}页 有“［…］”待汇报人填写的字段')
    rows.append((no, lay, len(sp['numbers']), ','.join(map(str, bc)) or '—',
                 min(bs) if bs else '—', min(fs) if fs else '—', 'Y' if sp['synthetic'] else ''))

# 9 溢出
if os.path.exists(FIT):
    fl = json.load(open(FIT, encoding='utf-8'))
    fails += [f"文字溢出（拟合日志）：第{d['slide']}页 {d['box']}（需 {d['need_in']} 英寸，有 {d['have_in']} 英寸）"
              for d in fl if d.get('overflow')]
else:
    warns.append(f'未找到拟合日志 {FIT}')
if os.path.exists(REN):
    ri = json.load(open(REN, encoding='utf-8'))
    fails += [f"文字溢出（预览渲染）：第{x['slide']}页 {x.get('shape', '')} {x['issue']}" for x in ri]
else:
    warns.append(f'未找到渲染问题清单 {REN}；请先运行 quick_render.py 或 LibreOffice 导出后逐页查看')

with open('ppt_check_r3.md', 'w', encoding='utf-8') as f:
    f.write(f'# 第三轮 PPT 自检报告\n\n文件：{PPT}　页数：{n}　文献：{nref} 条\n\n')
    f.write(f'## 失败（{len(fails)}）\n\n' + ''.join(f'- {x}\n' for x in fails) + '\n')
    f.write(f'## 提醒（{len(warns)}）\n\n' + ''.join(f'- {x}\n' for x in warns) + '\n')
    f.write('## 逐页概况\n\n| 页 | 版式 | 数字键数 | 正文文献 | 正文最小字号 | 脚注最小字号 | 模拟 |\n|---|---|---|---|---|---|---|\n')
    f.write(''.join('| ' + ' | '.join(map(str, r)) + ' |\n' for r in rows))
print('失败', len(fails), '提醒', len(warns), '→ ppt_check_r3.md')
for x in fails[:20]:
    print('  ✗', x)
